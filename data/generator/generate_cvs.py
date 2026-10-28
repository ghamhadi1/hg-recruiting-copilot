"""
Generates the fake CV set:
  data/cvs/          60 files (PDF + Word + 1 scanned image-only PDF)
  data/hard_cases/   extra broken files to prove "unreadable CVs are flagged"
  data/ground_truth.json

Run:  python data/generator/generate_cvs.py
"""
import copy, json, os, random, shutil, sys

from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

sys.path.insert(0, os.path.dirname(__file__))
from personas import DUPLICATES, PERSONAS, UPDATES  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "cvs")
HARD = os.path.join(ROOT, "hard_cases")
random.seed(42)

HEADINGS = [
    dict(summary="PROFILE", exp="EXPERIENCE", skills="SKILLS", langs="LANGUAGES", edu="EDUCATION"),
    dict(summary="Summary", exp="Work History", skills="Core Skills", langs="Languages", edu="Education & Certifications"),
    dict(summary="About me", exp="Professional Background", skills="Toolbox", langs="Languages spoken", edu="Studies"),
]


def sections(p):
    """Return the CV as a list of (kind, text) lines, layout depends on style."""
    h = HEADINGS[p["style"]]
    lines = [("name", p["name"]), ("sub", p["headline"])]
    contact = f'{p["location"]}  |  {p["email"]}  |  {p["phone"]}'
    if p["style"] != 1:
        lines.append(("text", contact))
    if p.get("updated") and p["style"] == 0:
        lines.append(("small", f'Last updated: {p["updated"]}'))
    lines += [("h", h["summary"]), ("text", p["summary"]), ("h", h["exp"])]
    for title, company, dates, bullets in p["jobs"]:
        if p["style"] == 2:
            lines.append(("job", f"{dates} — {title}, {company}"))
        else:
            lines.append(("job", f"{title} — {company}"))
            lines.append(("small", dates))
        lines += [("bullet", b) for b in bullets]
    lines.append(("h", h["skills"]))
    if p["style"] == 2:
        lines.append(("text", ", ".join(p["skills"])))
    else:
        lines += [("bullet", s) for s in p["skills"]]
    lines += [("h", h["langs"]), ("text", " · ".join(p["langs"]))]
    lines += [("h", h["edu"]), ("text", p["edu"])]
    lines += [("text", c) for c in p["certs"]]
    if p["style"] == 1:
        lines += [("h", "Contact"), ("text", contact)]
    if p.get("updated") and p["style"] != 0:
        lines.append(("small", f'CV updated {p["updated"]}'))
    return lines


def write_pdf(p, path):
    ss = getSampleStyleSheet()
    st = {
        "name": ParagraphStyle("n", parent=ss["Title"], fontSize=20, alignment=0 if p["style"] else 1),
        "sub": ParagraphStyle("s", parent=ss["Heading3"], textColor="#555555"),
        "h": ParagraphStyle("h", parent=ss["Heading2"], spaceBefore=10),
        "job": ParagraphStyle("j", parent=ss["Heading4"]),
        "small": ParagraphStyle("sm", parent=ss["Italic"], fontSize=9),
        "text": ss["BodyText"],
        "bullet": ParagraphStyle("b", parent=ss["BodyText"], leftIndent=14, bulletIndent=4),
    }
    story = []
    for kind, text in sections(p):
        txt = text.replace("&", "&amp;")
        story.append(Paragraph(("• " if kind == "bullet" else "") + txt, st[kind]))
        if kind == "h":
            story.append(Spacer(1, 2))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            title=f'{p["name"]} CV', author=p["name"])
    doc.build(story)


def write_docx(p, path):
    d = Document()
    d.styles["Normal"].font.size = Pt(10.5)
    if p["style"] == 1:
        # messy layout: header info inside a 2-column table
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = f'{p["name"]}\n{p["headline"]}'
        t.cell(0, 1).text = f'{p["location"]}\n{p["email"]}\n{p["phone"]}'
        rest = [l for l in sections(p) if l[0] not in ("name", "sub")]
        rest = [l for l in rest if not (l[0] == "text" and p["email"] in l[1])]
        rest = rest[:-2] if rest[-2:] and rest[-2][1] == "Contact" else rest
    else:
        rest = sections(p)
    for kind, text in rest:
        if kind == "name":
            d.add_heading(text, 0)
        elif kind in ("h",):
            d.add_heading(text.upper() if p["style"] == 2 else text, 1)
        elif kind == "job":
            d.add_heading(text, 3)
        elif kind == "bullet":
            d.add_paragraph(text, style="List Bullet")
        else:
            para = d.add_paragraph(text)
            if kind in ("small", "sub"):
                para.runs[0].italic = True
    d.core_properties.author = p["name"]
    d.save(path)


def write_scan(p, path):
    """Image-only PDF: no text layer, has to be OCR'd."""
    W, H = 1275, 1650  # 150 dpi letter
    img = Image.new("L", (W, H), 246)
    dr = ImageDraw.Draw(img)
    def font(size):
        for f in ("arial.ttf", "C:/Windows/Fonts/arial.ttf", "DejaVuSans.ttf"):
            try:
                return ImageFont.truetype(f, size)
            except OSError:
                continue
        return ImageFont.load_default()
    sizes = {"name": 40, "sub": 26, "h": 28, "job": 24, "small": 19, "text": 21, "bullet": 21}
    y = 70
    for kind, text in sections(p):
        f = font(sizes[kind])
        prefix = "- " if kind == "bullet" else ""
        words, line = (prefix + text).split(), ""
        for w in words:  # naive wrap
            test = (line + " " + w).strip()
            if dr.textlength(test, font=f) > W - 180:
                dr.text((90, y), line, fill=25, font=f); y += sizes[kind] + 8; line = w
            else:
                line = test
        dr.text((90, y), line, fill=25, font=f)
        y += sizes[kind] + (18 if kind == "h" else 10)
    # scanner look: slight rotation, blur, speckles
    img = img.rotate(0.8, fillcolor=235, resample=Image.BICUBIC)
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    px = img.load()
    for _ in range(9000):
        x, yy = random.randrange(W), random.randrange(H)
        px[x, yy] = random.randrange(120, 220)
    img.save(path, "PDF", resolution=150)


def render(p, path):
    {"pdf": write_pdf, "docx": write_docx, "scan": write_scan}[p["fmt"]](p, path)


def main():
    for d in (OUT, HARD):
        shutil.rmtree(d, ignore_errors=True)
        os.makedirs(d)
    by_key = {p["key"]: p for p in PERSONAS}
    truth = {}
    for p in PERSONAS:
        render(p, os.path.join(OUT, p["file"]))
        truth[p["key"]] = dict(name=p["name"], email=p["email"], files=[p["file"]],
                               active_file=p["file"], languages=p["langs"], skills=p["skills"],
                               scanned=p["fmt"] == "scan")
    for dup in DUPLICATES:
        src = by_key[dup["key"]]
        if not dup["changes"] and dup["fmt"] == src["fmt"]:
            shutil.copyfile(os.path.join(OUT, src["file"]), os.path.join(OUT, dup["file"]))
        else:
            q = copy.deepcopy(src); q.update(dup["changes"]); q.update(fmt=dup["fmt"], style=dup["style"])
            render(q, os.path.join(OUT, dup["file"]))
        truth[dup["key"]]["files"].append(dup["file"])
        truth[dup["key"]]["duplicate"] = True
    for up in UPDATES:
        q = copy.deepcopy(by_key[up["key"]]); q.update(up["changes"]); q.update(fmt=up["fmt"], style=up["style"])
        render(q, os.path.join(OUT, up["file"]))
        t = truth[up["key"]]
        t["files"].append(up["file"]); t["old_file"] = t["active_file"]; t["active_file"] = up["file"]
        t["skills"] = q["skills"]; t["updated"] = True

    # hard cases (not part of the 60)
    with open(os.path.join(HARD, "corrupt_cv.pdf"), "wb") as f:
        f.write(b"%PDF-1.4\n" + bytes(random.randrange(256) for _ in range(4000)))
    blank = Image.new("L", (1275, 1650), 250)
    ImageDraw.Draw(blank).ellipse((600, 800, 640, 830), fill=200)  # a coffee stain, nothing else
    blank.save(os.path.join(HARD, "blank_scan.pdf"), "PDF", resolution=150)

    with open(os.path.join(ROOT, "ground_truth.json"), "w", encoding="utf-8") as f:
        json.dump(truth, f, indent=2, ensure_ascii=False)

    files = os.listdir(OUT)
    print(f"{len(files)} CV files written ({sum(x.endswith('.pdf') for x in files)} PDF, "
          f"{sum(x.endswith('.docx') for x in files)} DOCX); {len(truth)} unique people")


if __name__ == "__main__":
    main()
