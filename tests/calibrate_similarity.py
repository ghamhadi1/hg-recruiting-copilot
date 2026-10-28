"""
Emulates Postgres pg_trgm similarity() on real CV pairs, to pick the threshold used in hg_ingest_cv
for "same person + same date: is it the same CV (duplicate) or a different one (replace)?".
Run: python tests/calibrate_similarity.py
"""
import os
import re

import docx
from pypdf import PdfReader

CVS = os.path.join(os.path.dirname(__file__), "..", "data", "cvs")


def text(name):
    p = os.path.join(CVS, name)
    if p.endswith(".pdf"):
        return "\n".join(pg.extract_text() or "" for pg in PdfReader(p).pages)
    d = docx.Document(p)
    parts = [x.text for x in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def trigrams(s):  # same rules as pg_trgm: words of alphanumerics, padded "  w "
    out = set()
    for w in re.findall(r"[a-z0-9]+", s.lower()):
        w = "  " + w + " "
        out.update(w[i:i + 3] for i in range(len(w) - 2))
    return out


def sim(a, b):
    A, B = trigrams(text(a)), trigrams(text(b))
    return len(A & B) / len(A | B)


PAIRS = [
    ("Omar_Haddad_CV.pdf", "omar haddad cv final.docx", "same CV, other format"),
    ("Bilal Qureshi.docx", "CV_BilalQureshi_FullStack.pdf", "same CV, other format"),
    ("Lina_Khoury_CV.docx", "Lina Khoury - Resume 2026.pdf", "UPDATED CV"),
    ("Layla_Hassan_CV.pdf", "Layla Hassan CV - updated Oct 2026.docx", "UPDATED CV"),
    ("Omar_Haddad_CV.pdf", "karim mansour resume.pdf", "different people"),
]
for a, b, label in PAIRS:
    print(f"{sim(a, b):.2f}  {label:<22} {a}  vs  {b}")
